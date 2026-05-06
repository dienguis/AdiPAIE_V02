"""
Addendum CODIR - V1.2 + V1.3 (mai 2026)
Genere un PPTX avec les evolutions depuis le PPT initial pour le CODIR.

Charte ELTON officielle (mars 2024) :
  - Bleu  : #1A73B5
  - Rouge : #E5003D
  - Navy  : #142E4D (titres slide)
  - Orange: #F18A1C (accents)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Charte ELTON
NAVY = RGBColor(0x14, 0x2E, 0x4D)
BLUE = RGBColor(0x1A, 0x73, 0xB5)
RED = RGBColor(0xE5, 0x00, 0x3D)
ORANGE = RGBColor(0xF1, 0x8A, 0x1C)
GREEN = RGBColor(0x2E, 0xA7, 0x5B)
PURPLE = RGBColor(0x7B, 0x5E, 0xA7)
GRAY_DARK = RGBColor(0x5C, 0x66, 0x79)
GRAY_LIGHT = RGBColor(0xF6, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height

BLANK = prs.slide_layouts[6]


def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    return s


def title_bar(s, title, subtitle=""):
    """Bandeau titre Navy + Orange en haut de la slide."""
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    underline = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.0), SW, Inches(0.05))
    underline.fill.solid()
    underline.fill.fore_color.rgb = ORANGE
    underline.line.fill.background()

    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.15), SW - Inches(0.8), Inches(0.85))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.05)

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(28)
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.name = "Calibri"

    if subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(14)
        r2.font.color.rgb = ORANGE
        r2.font.name = "Calibri"


def text_box(s, x, y, w, h, text, size=14, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return tb


def card(s, x, y, w, h, title, body, accent=BLUE, body_size=12):
    """Card blanche avec bordure colorée a gauche."""
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
    box.line.width = Pt(0.75)

    accent_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), h)
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()

    tb = s.shapes.add_textbox(x + Inches(0.2), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = NAVY
    r1.font.name = "Calibri"

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    p2.space_before = Pt(4)
    r2 = p2.add_run()
    r2.text = body
    r2.font.size = Pt(body_size)
    r2.font.color.rgb = GRAY_DARK
    r2.font.name = "Calibri"


def kpi_block(s, x, y, w, h, label, value, accent=ORANGE, sub=""):
    """Bloc KPI : label haut, gros chiffre, sous-titre."""
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
    top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.08))
    top.fill.solid()
    top.fill.fore_color.rgb = accent
    top.line.fill.background()

    tb = s.shapes.add_textbox(x + Inches(0.1), y + Inches(0.15), w - Inches(0.2), h - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = label.upper()
    r1.font.size = Pt(9)
    r1.font.bold = True
    r1.font.color.rgb = GRAY_DARK
    r1.font.name = "Calibri"

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = value
    r2.font.size = Pt(22)
    r2.font.bold = True
    r2.font.color.rgb = NAVY
    r2.font.name = "Calibri"

    if sub:
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        r3 = p3.add_run()
        r3.text = sub
        r3.font.size = Pt(9)
        r3.font.color.rgb = GRAY_DARK
        r3.font.name = "Calibri"


def footer(s, num):
    f = s.shapes.add_textbox(Inches(0.4), SH - Inches(0.4), SW - Inches(0.8), Inches(0.3))
    tf = f.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = f"SunuPaie ELTON | Addendum CODIR | Slide {num}"
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY_DARK
    r.font.italic = True
    r.font.name = "Calibri"


# =========================================================================
# SLIDE 1 - Titre
# =========================================================================
s = add_slide()
# Bandeau navy plein
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()

# Bandeau orange 5%
ob = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.8), SW, Inches(0.08))
ob.fill.solid()
ob.fill.fore_color.rgb = ORANGE
ob.line.fill.background()

text_box(s, Inches(0.6), Inches(1.3), SW - Inches(1.2), Inches(0.5),
         "ELTON Oil Company - SunuPaie", size=18, color=ORANGE, bold=True)

text_box(s, Inches(0.6), Inches(2.0), SW - Inches(1.2), Inches(0.8),
         "Addendum CODIR - V1.2 + V1.3", size=44, color=WHITE, bold=True)

text_box(s, Inches(0.6), Inches(3.1), SW - Inches(1.2), Inches(0.5),
         "Pilotage strategique DAF + DRH + Module Cout Reel Interimaires",
         size=18, color=WHITE)

text_box(s, Inches(0.6), Inches(4.0), SW - Inches(1.2), Inches(0.4),
         "Mai 2026 | Suite du PowerPoint initial (V1.0 + V1.1)",
         size=14, color=ORANGE, bold=False)

# Big numbers
kpi_block(s, Inches(1.0), Inches(5.0), Inches(2.6), Inches(1.5),
          "Dashboards initiaux", "6", ORANGE, "V1.0")
kpi_block(s, Inches(3.9), Inches(5.0), Inches(2.6), Inches(1.5),
          "Dashboards V1.2", "+4", BLUE, "DAF + DRH")
kpi_block(s, Inches(6.8), Inches(5.0), Inches(2.6), Inches(1.5),
          "Module V1.3", "+1", RED, "Cout Reel Interim")
kpi_block(s, Inches(9.7), Inches(5.0), Inches(2.6), Inches(1.5),
          "TOTAL", "11", GREEN, "Dashboards de pilotage")


# =========================================================================
# SLIDE 2 - Vue d'ensemble : 6 -> 11 dashboards
# =========================================================================
s = add_slide()
title_bar(s, "Evolution du module Tableaux de bord",
          "De 6 dashboards V1.0 vers 11 dashboards de pilotage strategique")

# Section V1.0 (6 dashboards a gauche)
text_box(s, Inches(0.4), Inches(1.3), Inches(4.3), Inches(0.4),
         "V1.0 - Bilan social classique (6 dashboards)", size=14, color=BLUE, bold=True)

dash_v10 = [
    ("N1 - Effectif detaille", "Pyramide ages H/F par categorie"),
    ("N2 - Analyse Effectif", "Global / Moyen / Evolution 8 ans"),
    ("N3 - Mouvements", "Arrivees / Departs par mois"),
    ("N4 - Remuneration", "Egalite salariale H/F"),
    ("N5 - Suivi Absences", "Conges, arrets, top 10"),
    ("N6 - Bilan Social Mensuel", "12 mois x 17 indicateurs"),
]
for i, (title, desc) in enumerate(dash_v10):
    y = Inches(1.85 + i * 0.65)
    card(s, Inches(0.4), y, Inches(4.3), Inches(0.55), title, desc, BLUE, body_size=10)

# Fleche
arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.85), Inches(3.3), Inches(0.55), Inches(0.6))
arr.fill.solid()
arr.fill.fore_color.rgb = ORANGE
arr.line.fill.background()

# Section V1.2 (4 dashboards au milieu)
text_box(s, Inches(5.55), Inches(1.3), Inches(3.85), Inches(0.4),
         "V1.2 - Pilotage DAF + DRH (+4)", size=14, color=ORANGE, bold=True)

dash_v12 = [
    ("N7 - Budget vs Realise", "DAF | Annuel + 5 ans glissants"),
    ("N8 - Provisions Sociales", "DAF | IDR (CCI Senegal) + Conges"),
    ("N9 - Cout Complet salarie", "Fully Loaded Cost + multiplicateur"),
    ("N10 - Conformite Senegal", "9 indicateurs reglementaires"),
]
for i, (title, desc) in enumerate(dash_v12):
    y = Inches(1.85 + i * 0.85)
    card(s, Inches(5.55), y, Inches(3.85), Inches(0.75), title, desc, ORANGE, body_size=10)

# Fleche
arr2 = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(9.55), Inches(3.3), Inches(0.55), Inches(0.6))
arr2.fill.solid()
arr2.fill.fore_color.rgb = RED
arr2.line.fill.background()

# Section V1.3 (1 dashboard a droite)
text_box(s, Inches(10.25), Inches(1.3), Inches(2.7), Inches(0.4),
         "V1.3 - Pilotage Interim (+1)", size=14, color=RED, bold=True)

card(s, Inches(10.25), Inches(1.85), Inches(2.7), Inches(2.0),
     "N11 - Cout Reel Interim",
     "DAF | TTC factures sociales d'interim vs cout theorique contrat",
     RED, body_size=10)

text_box(s, Inches(10.25), Inches(4.1), Inches(2.7), Inches(0.4),
         "Multiplicateur Brut->TTC ~1.91", size=11, color=NAVY, bold=True)
text_box(s, Inches(10.25), Inches(4.45), Inches(2.7), Inches(0.4),
         "Detection ecarts contrat/reel", size=11, color=GRAY_DARK)
text_box(s, Inches(10.25), Inches(4.8), Inches(2.7), Inches(0.4),
         "Validation factures avant paiement", size=11, color=GRAY_DARK)

# Bandeau bas synthese
b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(6.3), SW - Inches(0.8), Inches(0.7))
b.fill.solid()
b.fill.fore_color.rgb = NAVY
b.line.fill.background()
text_box(s, Inches(0.5), Inches(6.4), SW - Inches(1.0), Inches(0.5),
         "Investissement V1.2 + V1.3 : ~28 jours dev | Couvrent 100% des besoins DAF + DRH evoques au CODIR",
         size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

footer(s, 2)


# =========================================================================
# SLIDE 3 - V1.2 detail des 4 dashboards
# =========================================================================
s = add_slide()
title_bar(s, "V1.2 - 4 dashboards Pilotage DAF + DRH",
          "Reponse aux gaps identifies en revue strategique post-CODIR")

# 2x2 grid de cards detaillees
cards_v12 = [
    {
        "title": "TABLEAU N7 - BUDGET vs REALISE",
        "subtitle": "Pilotage DAF | Vue annuelle simplifiee",
        "kpis": ["Budget brut annuel", "Realise brut YTD", "Ecart valeur + %", "Sites en depassement"],
        "extras": ["Evolution 5 ans glissants", "Ventilation par site"],
        "color": BLUE
    },
    {
        "title": "TABLEAU N8 - PROVISIONS SOCIALES",
        "subtitle": "DAF + Audit | Conformite SYSCOHADA",
        "kpis": ["Provision IDR (bareme CCI Senegal)", "Provision Conges Payes", "Total a provisionner", "Variation mois"],
        "extras": ["Top 10 plus fortes provisions", "Ventilation par tranche anciennete"],
        "color": ORANGE
    },
    {
        "title": "TABLEAU N9 - COUT COMPLET salarie",
        "subtitle": "DAF + DRH | Fully Loaded Cost",
        "kpis": ["Cout total annuel", "Cout moyen / salarie", "Taux charges patronales", "Multiplicateur Net->Total"],
        "extras": ["Decomposition: Net/Cotis/Charges/Avantages", "Top 10 couts les plus eleves"],
        "color": PURPLE
    },
    {
        "title": "TABLEAU N10 - CONFORMITE SENEGAL",
        "subtitle": "DAF + DRH + Audit | Audit-ready",
        "kpis": ["SMIG respecte", "CDD < 2 ans", "Stages < 6 mois", "Soldes conges < 30 j"],
        "extras": ["9 indicateurs feux tricolores", "Score conformite global"],
        "color": GREEN
    },
]

positions = [
    (Inches(0.4), Inches(1.4)),
    (Inches(6.85), Inches(1.4)),
    (Inches(0.4), Inches(4.45)),
    (Inches(6.85), Inches(4.45)),
]
card_w = Inches(6.1)
card_h = Inches(2.95)

for (x, y), info in zip(positions, cards_v12):
    # Box
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, card_w, card_h)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
    box.line.width = Pt(0.75)

    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, card_w, Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = info["color"]
    accent.line.fill.background()

    # Titre
    text_box(s, x + Inches(0.2), y + Inches(0.18), card_w - Inches(0.3), Inches(0.4),
             info["title"], size=14, color=NAVY, bold=True)
    text_box(s, x + Inches(0.2), y + Inches(0.55), card_w - Inches(0.3), Inches(0.3),
             info["subtitle"], size=11, color=info["color"], bold=True)

    # KPIs (4 lignes)
    text_box(s, x + Inches(0.2), y + Inches(0.95), card_w - Inches(0.3), Inches(0.3),
             "KPIs cles :", size=10, color=GRAY_DARK, bold=True)
    for i, kpi in enumerate(info["kpis"]):
        text_box(s, x + Inches(0.3), y + Inches(1.2 + i * 0.22), card_w - Inches(0.4), Inches(0.25),
                 f"-  {kpi}", size=10, color=NAVY)

    # Extras
    text_box(s, x + Inches(0.2), y + Inches(2.15), card_w - Inches(0.3), Inches(0.3),
             "Sections :", size=10, color=GRAY_DARK, bold=True)
    for i, extra in enumerate(info["extras"]):
        text_box(s, x + Inches(0.3), y + Inches(2.4 + i * 0.22), card_w - Inches(0.4), Inches(0.25),
                 f"-  {extra}", size=10, color=NAVY)

footer(s, 3)


# =========================================================================
# SLIDE 4 - Module Budget RH (V1.2 entite socle)
# =========================================================================
s = add_slide()
title_bar(s, "Module Budget RH - Entite socle V1.2.1",
          "Saisie annuelle simplifiee suite revue DAF (mai 2026)")

# Avant / Apres
text_box(s, Inches(0.4), Inches(1.3), Inches(6.0), Inches(0.4),
         "AVANT (V1.2 initiale)", size=16, color=GRAY_DARK, bold=True)

text_box(s, Inches(0.4), Inches(1.85), Inches(6.0), Inches(0.4),
         "Saisie mensuelle x 9 rubriques x site", size=13, color=NAVY, bold=True)
text_box(s, Inches(0.4), Inches(2.25), Inches(6.0), Inches(0.4),
         "Volume : ~180 lignes par annee", size=12, color=GRAY_DARK)
text_box(s, Inches(0.4), Inches(2.65), Inches(6.0), Inches(0.4),
         "Complexite : mensualisation auto, gratifications S1/S2", size=12, color=GRAY_DARK)
text_box(s, Inches(0.4), Inches(3.05), Inches(6.0), Inches(0.4),
         "Refus DAF : trop complexe pour la realite metier", size=12, color=RED, bold=True)

# Fleche
arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.5), Inches(2.3), Inches(0.4), Inches(0.4))
arr.fill.solid()
arr.fill.fore_color.rgb = ORANGE
arr.line.fill.background()

text_box(s, Inches(7.0), Inches(1.3), Inches(6.0), Inches(0.4),
         "APRES (V1.2.1 - mise en prod)", size=16, color=GREEN, bold=True)

text_box(s, Inches(7.0), Inches(1.85), Inches(6.0), Inches(0.4),
         "Saisie annuelle sur le BRUT x site", size=13, color=NAVY, bold=True)
text_box(s, Inches(7.0), Inches(2.25), Inches(6.0), Inches(0.4),
         "Volume : 1 a 4 lignes par annee", size=12, color=GRAY_DARK)
text_box(s, Inches(7.0), Inches(2.65), Inches(6.0), Inches(0.4),
         "Saisie : Annee + Site + MontantBrutAnnuel", size=12, color=GRAY_DARK)
text_box(s, Inches(7.0), Inches(3.05), Inches(6.0), Inches(0.4),
         "Approuve par DAF en revue 2026-05-05", size=12, color=GREEN, bold=True)

# Stats demo
b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(3.7), SW - Inches(0.8), Inches(2.2))
b.fill.solid()
b.fill.fore_color.rgb = GRAY_LIGHT
b.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)

text_box(s, Inches(0.6), Inches(3.85), SW - Inches(1.2), Inches(0.4),
         "Donnees demo seedees (3 annees glissantes)", size=14, color=NAVY, bold=True)

kpi_block(s, Inches(0.7), Inches(4.4), Inches(3.0), Inches(1.3),
          "Budget 2024", "1,38 Md FCFA", BLUE, "Annee de reference")
kpi_block(s, Inches(3.9), Inches(4.4), Inches(3.0), Inches(1.3),
          "Budget 2025", "1,45 Md FCFA", ORANGE, "+5% inflation")
kpi_block(s, Inches(7.1), Inches(4.4), Inches(3.0), Inches(1.3),
          "Budget 2026", "1,52 Md FCFA", GREEN, "+5% N+1")
kpi_block(s, Inches(10.3), Inches(4.4), Inches(2.6), Inches(1.3),
          "Total seed", "12 lignes", PURPLE, "1 global + 3 sites x 3 ans")

# Bandeau benefice
b2 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(6.2), SW - Inches(0.8), Inches(0.7))
b2.fill.solid()
b2.fill.fore_color.rgb = NAVY
b2.line.fill.background()
text_box(s, Inches(0.5), Inches(6.3), SW - Inches(1.0), Inches(0.5),
         "Benefice : DAF saisit son budget CODIR en moins de 2 minutes par annee. Ecart automatique vs realise bulletins.",
         size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

footer(s, 4)


# =========================================================================
# SLIDE 5 - V1.3 Sprint 1 : Cout Reel Interimaires
# =========================================================================
s = add_slide()
title_bar(s, "V1.3 Sprint 1 - Cout Reel Interimaires",
          "Module pilotage de la facturation des societes d'interim (livre OK 2026-05-05)")

# Texte intro
text_box(s, Inches(0.4), Inches(1.25), SW - Inches(0.8), Inches(0.45),
         "Probleme metier resolu : aujourd'hui le cout interim affiche dans l'app est theorique (TauxJournalier x 22). " +
         "Or ELTON paye en realite le TTC de la facture mensuelle societe d'interim (debours + charges + commission + TVA).",
         size=12, color=GRAY_DARK)

# 4 KPIs reels du test d'integration
text_box(s, Inches(0.4), Inches(2.0), SW - Inches(0.8), Inches(0.4),
         "Test d'integration valide (mars 2026 - 31 lignes)",
         size=14, color=NAVY, bold=True)

kpi_block(s, Inches(0.4), Inches(2.5), Inches(3.0), Inches(1.3),
          "Brut imposable", "4,08 M FCFA", BLUE, "Salaire reel des 30 interim.")
kpi_block(s, Inches(3.6), Inches(2.5), Inches(3.0), Inches(1.3),
          "TTC paye par ELTON", "7,78 M FCFA", ORANGE, "Cout reel mensuel facture")
kpi_block(s, Inches(6.8), Inches(2.5), Inches(3.0), Inches(1.3),
          "Multiplicateur", "x1.91", RED, "1 FCFA brut = 1.91 FCFA reel")
kpi_block(s, Inches(10.0), Inches(2.5), Inches(2.95), Inches(1.3),
          "Annualise (~12 mois)", "~93 M FCFA", PURPLE, "Cout total 30 interim/an")

# Workflow utilisateur
text_box(s, Inches(0.4), Inches(4.0), SW - Inches(0.8), Inches(0.4),
         "Workflow utilisateur (5 etapes)",
         size=14, color=NAVY, bold=True)

steps = [
    ("1", "Recevoir facture", "Excel mensuel societe d'interim", BLUE),
    ("2", "Wizard import", "Annee + Mois + Societe", ORANGE),
    ("3", "Upload + Preview", "Verifier 31 lignes parsees", GREEN),
    ("4", "Confirmer commit", "Cocher Ecraser si refaire", PURPLE),
    ("5", "Dashboard N11", "Voir KPIs + ecart contrat/reel", RED),
]
sw_step = Inches(2.45)
for i, (num, title, desc, color) in enumerate(steps):
    x = Inches(0.4 + i * 2.55)
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(4.5), sw_step, Inches(1.5))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = color
    box.line.width = Pt(2)

    text_box(s, x + Inches(0.1), Inches(4.6), sw_step - Inches(0.2), Inches(0.5),
             num, size=20, color=color, bold=True, align=PP_ALIGN.CENTER)
    text_box(s, x + Inches(0.1), Inches(5.1), sw_step - Inches(0.2), Inches(0.4),
             title, size=12, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text_box(s, x + Inches(0.1), Inches(5.5), sw_step - Inches(0.2), Inches(0.4),
             desc, size=10, color=GRAY_DARK, align=PP_ALIGN.CENTER)

# Bandeau bas
b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(6.25), SW - Inches(0.8), Inches(0.7))
b.fill.solid()
b.fill.fore_color.rgb = ORANGE
b.line.fill.background()
text_box(s, Inches(0.5), Inches(6.35), SW - Inches(1.0), Inches(0.5),
         "Argument differenciant fort vs Fafadie Paie : aucune autre application au Senegal ne fait cette comparaison contrat / reel.",
         size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

footer(s, 5)


# =========================================================================
# SLIDE 6 - Roadmap V1.3 suite
# =========================================================================
s = add_slide()
title_bar(s, "Roadmap V1.3 - Suite des sprints",
          "4 modules complementaires apres validation CODIR")

text_box(s, Inches(0.4), Inches(1.3), SW - Inches(0.8), Inches(0.4),
         "Sprint 1 livre (mai 2026) : V1.3 - Cout Reel Interimaires (production-ready)",
         size=14, color=GREEN, bold=True)

# Sprints suivants
sprints = [
    {
        "num": "Sprint 2", "duree": "7 jours", "titre": "GPEC / Skill Matrix",
        "desc": "Cartographie competences par equipe, identification gaps, plan de formation cible. Necessite refonte modele Competences (entite absente aujourd'hui).",
        "color": BLUE
    },
    {
        "num": "Sprint 3", "duree": "5 jours", "titre": "Plan de releve (Succession Planning)",
        "desc": "Identifier postes critiques + 2 successeurs avec readiness. Workflow de validation. Reduit le risque continuite operationnelle.",
        "color": ORANGE
    },
    {
        "num": "Sprint 4", "duree": "4 jours", "titre": "Pay Equity Gap H/F",
        "desc": "Calcul ecart H/F par poste equivalent. Methodologie a definir avec DRH. Sensible juridiquement, prepare la conformite RSE.",
        "color": PURPLE
    },
    {
        "num": "Sprint 5", "duree": "4 jours", "titre": "Suivi entretiens annuels",
        "desc": "Workflow EntretienAnnuel complet : objectifs N+1, plan d'action, validation N+2. Entite existe mais workflow incomplet.",
        "color": GREEN
    },
]

for i, sp in enumerate(sprints):
    y = Inches(2.05 + i * 1.05)
    # Card
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), y, SW - Inches(0.8), Inches(0.95))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)

    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), y, Inches(0.1), Inches(0.95))
    accent.fill.solid()
    accent.fill.fore_color.rgb = sp["color"]
    accent.line.fill.background()

    # Sprint num + duree
    text_box(s, Inches(0.6), y + Inches(0.15), Inches(2.0), Inches(0.4),
             sp["num"], size=16, color=sp["color"], bold=True)
    text_box(s, Inches(0.6), y + Inches(0.5), Inches(2.0), Inches(0.4),
             sp["duree"], size=11, color=GRAY_DARK)

    # Titre
    text_box(s, Inches(2.7), y + Inches(0.15), Inches(3.3), Inches(0.4),
             sp["titre"], size=15, color=NAVY, bold=True)
    text_box(s, Inches(2.7), y + Inches(0.5), Inches(3.3), Inches(0.4),
             "Demarrage post-CODIR", size=10, color=GRAY_DARK)

    # Description
    text_box(s, Inches(6.2), y + Inches(0.1), SW - Inches(6.5), Inches(0.85),
             sp["desc"], size=11, color=GRAY_DARK)

# Bandeau total
b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(6.4), SW - Inches(0.8), Inches(0.55))
b.fill.solid()
b.fill.fore_color.rgb = NAVY
b.line.fill.background()
text_box(s, Inches(0.5), Inches(6.45), SW - Inches(1.0), Inches(0.45),
         "Total V1.3 Sprints 2-5 : ~20 jours dev (4-5 semaines) | Sous reserve d'arbitrage CODIR",
         size=12, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

footer(s, 6)


# =========================================================================
# SLIDE 7 - Decision CODIR
# =========================================================================
s = add_slide()
title_bar(s, "Decision CODIR demandee",
          "Validation V1.2 + V1.3 Sprint 1 + arbitrage suite")

# 3 grandes cards de decision
decisions = [
    {
        "num": "1",
        "title": "Validation mise en production V1.2 + V1.3 Sprint 1",
        "items": [
            "5 nouveaux dashboards de pilotage (N7-N11)",
            "Module Budget RH (saisie annuelle)",
            "Module Cout Reel Interimaires (test OK)",
            "Tous tests d'integration passes",
            "Documentation help complete"
        ],
        "color": GREEN
    },
    {
        "num": "2",
        "title": "Arbitrage roadmap V1.3 suite (Sprints 2-5)",
        "items": [
            "Sprint 2 - GPEC / Skill Matrix (7j)",
            "Sprint 3 - Plan de releve (5j)",
            "Sprint 4 - Pay Equity Gap H/F (4j)",
            "Sprint 5 - Suivi entretiens annuels (4j)",
            "Priorisation par DRH a fournir"
        ],
        "color": ORANGE
    },
    {
        "num": "3",
        "title": "Validation parametres metier",
        "items": [
            "Bareme IDR ELTON (CCI standard ou avenant ?)",
            "Seuils alerte budget vs realise (5% ?)",
            "Liste des sites prioritaires pilotage",
            "Liste des societes d'interim partenaires",
            "Date demarrage prochain sprint"
        ],
        "color": BLUE
    },
]

for i, d in enumerate(decisions):
    x = Inches(0.4 + i * 4.3)
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.4), Inches(4.1), Inches(4.6))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = d["color"]
    box.line.width = Pt(3)

    # Numero
    circle = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(1.55), Inches(1.55), Inches(1.0), Inches(1.0))
    circle.fill.solid()
    circle.fill.fore_color.rgb = d["color"]
    circle.line.fill.background()
    text_box(s, x + Inches(1.55), Inches(1.7), Inches(1.0), Inches(0.7),
             d["num"], size=36, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Titre
    text_box(s, x + Inches(0.2), Inches(2.7), Inches(3.7), Inches(0.6),
             d["title"], size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # Items
    for j, item in enumerate(d["items"]):
        text_box(s, x + Inches(0.3), Inches(3.4 + j * 0.4), Inches(3.6), Inches(0.4),
                 f"-  {item}", size=10, color=GRAY_DARK)

# Bandeau bas
b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(6.2), SW - Inches(0.8), Inches(0.85))
b.fill.solid()
b.fill.fore_color.rgb = RED
b.line.fill.background()
text_box(s, Inches(0.5), Inches(6.3), SW - Inches(1.0), Inches(0.7),
         "Merci - Questions / Discussion",
         size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

footer(s, 7)


# =========================================================================
# Sauvegarde
# =========================================================================
output = '/sessions/pensive-modest-edison/mnt/AdiPAIE_V02/SunuPaie_CODIR_Update_V12_V13.pptx'
prs.save(output)

import os
print(f"OK -> {output}")
print(f"Taille : {os.path.getsize(output):,} octets")
print(f"Slides : {len(prs.slides)}")
